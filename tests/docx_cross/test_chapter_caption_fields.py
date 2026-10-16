"""DOCX captions: chapter-scoped "Figure 6.3" / "Figure A.2" fields, or one running count."""

from docx import Document
from docx.oxml.ns import qn

from paradoc.io.word.captions import rebuild_caption


def _instructions(paragraph):
    return [el.text.strip() for el in paragraph._p.iter(qn("w:instrText"))]


def test_chapter_caption_main():
    doc = Document()
    p = doc.add_paragraph("old")
    rebuild_caption(p, "Figure", "A caption", is_appendix=False, should_restart=True)
    assert _instructions(p) == ['STYLEREF \\s "Heading 1" \\n \\t', "SEQ Figure \\* ARABIC \\r 1 \\s 1"]
    # Placeholder text until Word updates the fields
    assert p.text == "Figure 1.1: A caption"


def test_chapter_caption_appendix():
    doc = Document()
    p = doc.add_paragraph("old")
    rebuild_caption(p, "Table", "Extra", is_appendix=True, should_restart=False)
    assert _instructions(p) == ['STYLEREF \\s "Appendix" \\n \\t', "SEQ Table \\* ARABIC \\s 1"]
    assert p.text == "Table A.1: Extra"


def test_running_caption_without_chapters():
    doc = Document()
    p = doc.add_paragraph("old")
    rebuild_caption(p, "Figure", "Plain", is_appendix=True, should_restart=False, chapter_numbering=False)
    assert _instructions(p) == ["SEQ Figure \\* ARABIC"]
    assert p.text == "Figure 1: Plain"
