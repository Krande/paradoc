"""Export paths that must not fall over on an ordinary document or an ordinary machine."""

import pytest
from docx import Document
from docx.oxml.ns import qn

from paradoc.exceptions import LatexNotInstalled
from paradoc.io.pdf import exporter as pdf_exporter
from paradoc.io.word.reference_helper import ReferenceHelper
from paradoc.io.word.utils import request_field_update_on_open


def test_a_table_directly_after_another_has_no_caption_rather_than_crashing():
    doc = Document()
    doc.add_paragraph("Table 1-1: The first table")
    first = doc.add_table(rows=1, cols=1)
    doc.add_table(rows=1, cols=1)

    helper = ReferenceHelper()
    # The block before the second table is the first table, which has no `.text`.
    assert helper._extract_caption_text_from_caption(first) is None
    assert helper._extract_caption_text_from_caption(doc.paragraphs[0]) == "The first table"


def test_field_update_on_open_goes_where_the_schema_puts_it_and_can_be_removed(tmp_path):
    path = tmp_path / "doc.docx"
    Document().save(str(path))

    request_field_update_on_open(path)
    settings = Document(str(path)).settings.element
    flag = settings.find(qn("w:updateFields"))
    assert flag is not None and flag.get(qn("w:val")) == "true"
    # Word rejects settings.xml with elements out of schema order; `compat` follows `updateFields`.
    tags = [el.tag for el in settings]
    if qn("w:compat") in tags:
        assert tags.index(qn("w:updateFields")) < tags.index(qn("w:compat"))

    request_field_update_on_open(path)  # idempotent
    assert len(Document(str(path)).settings.element.findall(qn("w:updateFields"))) == 1

    request_field_update_on_open(path, request=False)
    assert Document(str(path)).settings.element.find(qn("w:updateFields")) is None


def test_pdf_engine_falls_back_past_a_missing_xelatex(monkeypatch):
    monkeypatch.delenv("PARADOC_PDF_ENGINE", raising=False)
    monkeypatch.setattr(pdf_exporter.shutil, "which", lambda name: "/bin/tectonic" if name == "tectonic" else None)
    assert pdf_exporter.resolve_pdf_engine() == "/bin/tectonic"


def test_pdf_engine_override_and_absence_are_reported(monkeypatch):
    monkeypatch.setattr(pdf_exporter.shutil, "which", lambda name: None)
    monkeypatch.setenv("PARADOC_PDF_ENGINE", "no-such-engine")
    with pytest.raises(LatexNotInstalled, match="no-such-engine"):
        pdf_exporter.resolve_pdf_engine()
    monkeypatch.delenv("PARADOC_PDF_ENGINE")
    with pytest.raises(LatexNotInstalled, match="tectonic"):
        pdf_exporter.resolve_pdf_engine()
