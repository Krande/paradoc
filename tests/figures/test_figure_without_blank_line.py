"""A `{#fig:...}` image written directly under a text line must still export as a figure (#20)."""

import shutil

from docx import Document

from paradoc import OneDoc
from paradoc.common import isolate_block_figures


def test_isolate_block_figures_inserts_blank_lines():
    md = "Text.\n![Cap](a.png){#fig:a}\nMore text.\n"
    assert isolate_block_figures(md) == "Text.\n\n![Cap](a.png){#fig:a}\n\nMore text.\n"


def test_isolate_block_figures_leaves_other_lines_alone():
    md = (
        "Text.\n\n![Cap](a.png){#fig:a}\n\n"  # already separated
        "Inline ![x](b.png) image.\n"  # unlabelled inline image
        "Text with ![Cap](c.png){#fig:c} mid-line.\n"  # not alone on its line
        "```\nText.\n![Cap](d.png){#fig:d}\n```\n"  # inside a fenced code block
    )
    assert isolate_block_figures(md) == md


def test_figure_without_blank_line_exported_to_docx(files_dir, tmp_path):
    src = tmp_path / "doc"
    (src / "00-main" / "images").mkdir(parents=True)
    shutil.copy(files_dir / "doc_lorum/01-app/images/qc1.png", src / "00-main/images/qc1.png")
    (src / "00-main/main.md").write_text(
        "# Chapter\n\n"
        "Text directly followed by a figure.\n"
        "![Figure without blank line](images/qc1.png){#fig:noblank}\n"
        "Text directly after the figure.\n",
        encoding="utf-8",
    )

    one = OneDoc(src, work_dir=tmp_path / "work")
    one.compile("fig_no_blank", auto_open=False, export_format="docx")

    doc = Document(str(next((tmp_path / "work").rglob("fig_no_blank.docx"))))
    assert sum(len(p._p.xpath(".//pic:pic")) for p in doc.paragraphs) == 1
    captions = [p.text for p in doc.paragraphs if p.style.name == "Image Caption"]
    assert len(captions) == 1 and "Figure without blank line" in captions[0]
